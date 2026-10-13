import argparse
import csv
import io
import json
import mimetypes
import os
import re
import shutil
import sys
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload, MediaIoBaseDownload, MediaIoBaseUpload
import openpyxl
from openpyxl.utils import column_index_from_string

SCOPES = ["https://www.googleapis.com/auth/drive"]
DRIVE_ID_ALIASES = {
    "1vL7G8rC1j-X0D4mK5sP9qW2eN6tB8yH1": "1idzoh4DyV8YWunTqROpSsIvJTyKV2vnL",
}


def resolve_drive_id(fid):
    return DRIVE_ID_ALIASES.get(fid, fid) if fid else fid

CONFIG_TOKEN_PATH = os.path.expanduser(r"~\.gemini\config\gdrive_tokens.json")
FALLBACK_TOKEN_PATH = r"H:\Antigravity\SAO Mockup\gdrive_tokens.json"
CONFIG_CREDS_PATH = os.path.expanduser(r"~\.gemini\config\drive_credentials.json")
FALLBACK_CREDS_PATH = r"H:\Antigravity\SAO Mockup\drive_credentials.json"


def get_token_path():
    if os.path.exists(CONFIG_TOKEN_PATH):
        return CONFIG_TOKEN_PATH
    if os.path.exists(FALLBACK_TOKEN_PATH):
        os.makedirs(os.path.dirname(CONFIG_TOKEN_PATH), exist_ok=True)
        shutil.copy2(FALLBACK_TOKEN_PATH, CONFIG_TOKEN_PATH)
        return CONFIG_TOKEN_PATH
    return CONFIG_TOKEN_PATH


def get_creds_path():
    if os.path.exists(CONFIG_CREDS_PATH):
        return CONFIG_CREDS_PATH
    if os.path.exists(FALLBACK_CREDS_PATH):
        os.makedirs(os.path.dirname(CONFIG_CREDS_PATH), exist_ok=True)
        shutil.copy2(FALLBACK_CREDS_PATH, CONFIG_CREDS_PATH)
        return CONFIG_CREDS_PATH
    return CONFIG_CREDS_PATH


def load_credentials(interactive=False):
    token_path = get_token_path()
    creds = None
    if os.path.exists(token_path):
        creds = Credentials.from_authorized_user_file(token_path, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
            with open(token_path, "w", encoding="utf-8") as f:
                f.write(creds.to_json())
        elif interactive:
            creds_path = get_creds_path()
            flow = InstalledAppFlow.from_client_secrets_file(creds_path, SCOPES)
            creds = flow.run_local_server(port=0)
            with open(token_path, "w", encoding="utf-8") as f:
                f.write(creds.to_json())
        else:
            return None
    return creds


def get_drive_service():
    creds = load_credentials(interactive=False)
    if not creds:
        print("Error: Not logged in. Run `gdrive auth login` or `gsheets auth login`.", file=sys.stderr)
        sys.exit(1)
    return build("drive", "v3", credentials=creds)


def handle_auth(tool_name, args):
    action = args[0] if args else "status"
    if action == "status":
        creds = load_credentials(interactive=False)
        if not creds:
            print(f"{tool_name}: Not logged in")
            sys.exit(1)
        drive = build("drive", "v3", credentials=creds)
        about = drive.about().get(fields="user").execute()
        user = about.get("user", {})
        print(f"{tool_name}: Logged in as {user.get('displayName')} <{user.get('emailAddress')}>")
    elif action == "login":
        creds = load_credentials(interactive=True)
        drive = build("drive", "v3", credentials=creds)
        about = drive.about().get(fields="user").execute()
        user = about.get("user", {})
        print(f"{tool_name}: Successfully logged in as {user.get('displayName')} <{user.get('emailAddress')}>")
    else:
        print(f"Unknown auth action: {action}", file=sys.stderr)
        sys.exit(1)


def guess_mime(path):
    if path.endswith(".md"):
        return "text/markdown"
    if path.endswith(".json"):
        return "application/json"
    if path.endswith(".py"):
        return "text/x-python"
    mime, _ = mimetypes.guess_type(path)
    return mime or "application/octet-stream"


def handle_gdrive(argv):
    if not argv or argv[0] in ("-h", "--help"):
        print("Usage: gdrive [auth|readonly|mutate] <subcommand> [options]")
        return
    ns = argv[0]
    rest = argv[1:]
    if ns == "auth":
        handle_auth("gdrive", rest)
        return

    drive = get_drive_service()
    if ns == "readonly":
        sub = rest[0] if rest else "list"
        sub_args = rest[1:]
        if sub == "list":
            parser = argparse.ArgumentParser()
            parser.add_argument("--parent", default=None)
            parser.add_argument("--query", "-q", default=None)
            parser.add_argument("--max", type=int, default=100)
            opts = parser.parse_args(sub_args)
            opts.parent = resolve_drive_id(opts.parent)
            q_parts = ["trashed=false"]
            if opts.parent:
                q_parts.append(f"'{opts.parent}' in parents")
            if opts.query:
                q_parts.append(f"({opts.query})")
            q = " and ".join(q_parts)
            res = drive.files().list(
                q=q,
                pageSize=opts.max,
                fields="files(id, name, mimeType, modifiedTime, size)"
            ).execute()
            files = res.get("files", [])
            print(json.dumps({"status": "success", "count": len(files), "files": files}, indent=2))
        elif sub == "get":
            file_id = resolve_drive_id(sub_args[0])
            meta = drive.files().get(fileId=file_id, fields="id, name, mimeType, modifiedTime, size, parents").execute()
            print(json.dumps(meta, indent=2))
        elif sub == "download":
            file_id = resolve_drive_id(sub_args[0])
            dest = sub_args[1] if len(sub_args) > 1 else None
            meta = drive.files().get(fileId=file_id, fields="id, name, mimeType").execute()
            if not dest:
                dest = meta["name"]
            req = drive.files().get_media(fileId=file_id)
            with io.FileIO(dest, "wb") as fh:
                downloader = MediaIoBaseDownload(fh, req)
                done = False
                while not done:
                    _, done = downloader.next_chunk()
            print(json.dumps({"status": "downloaded", "id": file_id, "path": os.path.abspath(dest)}))
        else:
            print(f"Unknown gdrive readonly subcommand: {sub}", file=sys.stderr)
            sys.exit(1)

    elif ns == "mutate":
        sub = rest[0] if rest else ""
        sub_args = rest[1:]
        if sub == "upload":
            parser = argparse.ArgumentParser()
            parser.add_argument("local_path")
            parser.add_argument("--parent", default=None)
            parser.add_argument("--name", default=None)
            opts = parser.parse_args(sub_args)
            opts.parent = resolve_drive_id(opts.parent)
            local_path = opts.local_path
            name = opts.name or os.path.basename(local_path)
            mime = guess_mime(local_path)
            media = MediaFileUpload(local_path, mimetype=mime, resumable=True)
            # Check if file with same name already exists in parent to update in-place
            existing_id = None
            if opts.parent:
                q = f"'{opts.parent}' in parents and name='{name}' and trashed=false"
                res = drive.files().list(q=q, fields="files(id, name)").execute()
                items = res.get("files", [])
                if items:
                    existing_id = items[0]["id"]
            if existing_id:
                updated = drive.files().update(
                    fileId=existing_id,
                    media_body=media,
                    fields="id, name, mimeType, modifiedTime"
                ).execute()
                updated["action"] = "updated_existing"
                print(json.dumps(updated, indent=2))
            else:
                body = {"name": name}
                if opts.parent:
                    body["parents"] = [opts.parent]
                created = drive.files().create(
                    body=body,
                    media_body=media,
                    fields="id, name, mimeType, modifiedTime"
                ).execute()
                created["action"] = "created_new"
                print(json.dumps(created, indent=2))
        elif sub == "update":
            file_id = resolve_drive_id(sub_args[0])
            local_path = sub_args[1]
            mime = guess_mime(local_path)
            media = MediaFileUpload(local_path, mimetype=mime, resumable=True)
            updated = drive.files().update(
                fileId=file_id,
                media_body=media,
                fields="id, name, mimeType, modifiedTime"
            ).execute()
            print(json.dumps(updated, indent=2))
        elif sub == "mkdir":
            parser = argparse.ArgumentParser()
            parser.add_argument("name")
            parser.add_argument("--parent", default=None)
            opts = parser.parse_args(sub_args)
            opts.parent = resolve_drive_id(opts.parent)
            q_parts = [
                f"name='{opts.name}'",
                "mimeType='application/vnd.google-apps.folder'",
                "trashed=false"
            ]
            if opts.parent:
                q_parts.append(f"'{opts.parent}' in parents")
            res = drive.files().list(q=" and ".join(q_parts), fields="files(id, name)").execute()
            items = res.get("files", [])
            if items:
                print(json.dumps({"status": "existing", "id": items[0]["id"], "name": items[0]["name"]}, indent=2))
            else:
                body = {"name": opts.name, "mimeType": "application/vnd.google-apps.folder"}
                if opts.parent:
                    body["parents"] = [opts.parent]
                folder = drive.files().create(body=body, fields="id, name").execute()
                print(json.dumps({"status": "created", "id": folder["id"], "name": folder["name"]}, indent=2))
        else:
            print(f"Unknown gdrive mutate subcommand: {sub}", file=sys.stderr)
            sys.exit(1)


def parse_a1_range(range_str, wb):
    if "!" in range_str:
        sheet_name, cell_range = range_str.split("!", 1)
        sheet_name = sheet_name.strip("'")
    elif range_str in wb.sheetnames:
        return wb[range_str], None
    else:
        sheet_name = wb.sheetnames[0]
        cell_range = range_str
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
    else:
        ws = wb[sheet_name]
    return ws, cell_range


def parse_cell_coord(cell_ref):
    m = re.match(r"^([A-Za-z]+)(\d+)$", cell_ref.strip())
    if not m:
        return 1, 1
    col = column_index_from_string(m.group(1).upper())
    row = int(m.group(2))
    return row, col


def export_workbook(drive, spreadsheet_id):
    xlsx_bytes = drive.files().export(
        fileId=spreadsheet_id,
        mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ).execute()
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def save_workbook(drive, spreadsheet_id, wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    media = MediaIoBaseUpload(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        resumable=False
    )
    return drive.files().update(
        fileId=spreadsheet_id,
        media_body=media,
        fields="id, name, modifiedTime"
    ).execute()


def handle_gsheets(argv):
    if not argv or argv[0] in ("-h", "--help"):
        print("Usage: gsheets [auth|readonly|mutate] <subcommand> [options]")
        return
    ns = argv[0]
    rest = argv[1:]
    if ns == "auth":
        handle_auth("gsheets", rest)
        return

    drive = get_drive_service()
    if ns == "readonly":
        sub = rest[0] if rest else "read"
        sub_args = rest[1:]
        if sub == "read":
            spreadsheet_id = sub_args[0]
            range_str = sub_args[1] if len(sub_args) > 1 else "A1:Z100"
            wb = export_workbook(drive, spreadsheet_id)
            ws, cell_range = parse_a1_range(range_str, wb)
            rows_out = []
            if cell_range and ":" in cell_range:
                for row in ws[cell_range]:
                    vals = [("" if c.value is None else str(c.value)) for c in row]
                    if any(vals):
                        rows_out.append(vals)
            elif cell_range:
                c = ws[cell_range]
                rows_out.append(["" if c.value is None else str(c.value)])
            else:
                for row in ws.iter_rows(values_only=True):
                    vals = [("" if v is None else str(v)) for v in row]
                    if any(vals):
                        rows_out.append(vals)
            print(json.dumps({"sheet": ws.title, "range": range_str, "rows": rows_out}, indent=2))
        else:
            print(f"Unknown gsheets readonly subcommand: {sub}", file=sys.stderr)
            sys.exit(1)

    elif ns == "mutate":
        sub = rest[0] if rest else ""
        sub_args = rest[1:]
        if sub == "write":
            spreadsheet_id = sub_args[0]
            range_str = sub_args[1]
            raw_vals = sub_args[2:]
            wb = export_workbook(drive, spreadsheet_id)
            ws, cell_range = parse_a1_range(range_str, wb)
            start_cell = cell_range.split(":")[0] if cell_range else "A1"
            start_row, start_col = parse_cell_coord(start_cell)
            if len(raw_vals) == 1 and ("," in raw_vals[0] or "\t" in raw_vals[0]) and ":" in (cell_range or ""):
                parsed_rows = list(csv.reader([raw_vals[0]]))
            else:
                parsed_rows = [raw_vals]
            for r_idx, row_vals in enumerate(parsed_rows):
                for c_idx, val in enumerate(row_vals):
                    ws.cell(row=start_row + r_idx, column=start_col + c_idx, value=val)
            res = save_workbook(drive, spreadsheet_id, wb)
            print(json.dumps({"status": "updated", "sheet": ws.title, "range": range_str, "file": res}, indent=2))
        elif sub == "append":
            spreadsheet_id = sub_args[0]
            sheet_or_range = sub_args[1]
            raw_vals = sub_args[2:]
            wb = export_workbook(drive, spreadsheet_id)
            ws, _ = parse_a1_range(sheet_or_range, wb)
            if len(raw_vals) == 1:
                row_vals = next(csv.reader([raw_vals[0]]))
            else:
                row_vals = raw_vals
            ws.append(row_vals)
            res = save_workbook(drive, spreadsheet_id, wb)
            print(json.dumps({"status": "appended", "sheet": ws.title, "row": row_vals, "file": res}, indent=2))
        else:
            print(f"Unknown gsheets mutate subcommand: {sub}", file=sys.stderr)
            sys.exit(1)


if __name__ == "__main__":
    tool = os.path.basename(sys.argv[1]).lower() if len(sys.argv) > 1 else "gdrive"
    if tool == "gsheets":
        handle_gsheets(sys.argv[2:])
    else:
        handle_gdrive(sys.argv[2:])
